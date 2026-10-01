#!/usr/bin/env python3
"""
mcc_universal_mapper.py
------------------------
ONE script for every MCC NEET-UG counselling round PDF -- Round 1, Round 2,
Round 3, Round 4, Round 5, whichever comes next. You don't need a separate
script per round; just point this at whatever PDF MCC publishes.

WHY ONE SCRIPT WORKS FOR EVERY ROUND
    Round 1 PDFs use a "simple" layout: one allotment per candidate.
        SNo, Rank, Allotted Quota, Allotted Institute, Course,
        Alloted Category, Candidate Category, Remarks               (8 columns)

    Round 2 onward, MCC shows EVERY round so far side by side in one row --
    Round 1 | Round 2 | ... | Round N -- so the table gets one block wider
    every round:
        Rank,
        [Round 1]   Quota, Institute, Course, Remarks
        [Round 2]   Quota, Institute, Course, Remarks
        ...
        [Round N]   Quota, Institute, Course, Alloted Category,
                     Candidate Category, Option No., Remarks
    Only the LAST block (the latest round) has the extra category/option
    columns; every earlier block is just Quota/Institute/Course/Remarks.
    That means no matter how many round-blocks are present, the layout
    always ends the same way, so the latest round's own "Allotted
    Institute" is always the 6th-from-last column in the row. This script
    reads each PDF's own table structure (via pdfplumber's table
    detection, not hardcoded pixel positions) and picks that column
    automatically -- so it keeps working as the table grows for Round 4,
    Round 5, etc, without any changes needed.

WHICH INSTITUTE GETS MAPPED
    - Round 1 style PDF (8 columns): the one "Allotted Institute" column.
    - Round 2-onward style PDF: ONLY the latest round's own "Allotted
      Institute" column (the 6th-from-last). If it's blank/"-" (candidate
      did not opt for upgradation, did not fill fresh choices, or was not
      allotted this round), that row is marked "N/A" and left unmapped --
      by design, nothing is pulled from Round 1 or any earlier round to
      fill the gap. (Note: even if that were wanted, every earlier
      round's column in this PDF layout only ever contains a short
      institute name with no address, so it could never match
      FORMAT.xlsx anyway.)

MATCHING AGAINST FORMAT.xlsx
    Normalize by stripping ALL whitespace (not just collapsing it) and
    lower-casing. This reaches the practical matching ceiling for this
    data (verified on real Round 1 & Round 2 files: additionally
    stripping punctuation on top of this gained zero extra matches). Any
    institute still unmatched after that is either (a) genuinely absent
    from FORMAT.xlsx (e.g. Nursing colleges, if FORMAT.xlsx only lists
    MBBS/BDS) or (b) has a real wording difference in the source PDF
    itself -- neither of which normalization can fix. Those rows are
    still written out in full, just with blank mapped columns and Match
    Status = "No Match", so nothing is ever silently dropped.

USAGE
    python3 mcc_universal_mapper.py <input_pdf> [output_xlsx] [--format-xlsx PATH]

    FORMAT.xlsx path is fixed below (edit FORMAT_XLSX_PATH) -- override
    with --format-xlsx if you need to point at a different copy.
"""

import sys
import re
import argparse
from pathlib import Path
from collections import OrderedDict

import pdfplumber
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# FORMAT.xlsx is the fixed reference/lookup workbook -- it doesn't change
# between rounds, only the input PDF does.
# --------------------------------------------------------------------------
FORMAT_XLSX_PATH = r"C:\Users\Admin\Desktop\MCC FORMAT MAPPING\FORMAT.xlsx"

ROUND_LABEL_RE = re.compile(r"Round\s*\d+", re.IGNORECASE)


def normalize_institute(text):
    """Case-fold and strip ALL whitespace for robust matching.

    PDF line-wraps can land in slightly different spots between different
    occurrences of the exact same institute (extra/missing space around a
    hyphen, a comma, between words, etc). Institute names/addresses never
    legitimately depend on exact spacing, so we remove whitespace entirely
    rather than just collapsing it -- this reaches the practical matching
    ceiling for this data (verified against real Round 1 & Round 2 files;
    additionally stripping punctuation gained zero extra matches on top
    of this).
    """
    if text is None:
        return ""
    text = str(text).replace("\n", " ")
    text = re.sub(r"\s+", "", text)
    return text.lower()


TRAILING_STATE_PINCODE_RE = re.compile(r",\s*[A-Za-z .()&\-]+,\s*\d{4,6}\s*$")


def lookup_institute(inst, lookup):
    """
    Try an exact normalized match first. If that fails, some PDFs (seen
    on older/different-year reports) append a redundant trailing
    ", <State>, <Pincode>" that FORMAT.xlsx's address text doesn't have
    -- e.g. PDF: "...Kunraghat, Gorakhpur, Uttar Pradesh, 273008" vs
    FORMAT.xlsx: "...Kunraghat, Gorakhpur". Stripping that trailing
    suffix and retrying recovers a meaningful chunk of those rows. This
    is a fallback only -- a genuine year/format mismatch between the PDF
    and FORMAT.xlsx will still leave many rows unmatched, since the
    remaining differences aren't just this one suffix pattern.

    Returns (mapped_dict_or_None, matched_via) where matched_via is
    "direct", "suffix-stripped", or None.
    """
    key = normalize_institute(inst)
    mapped = lookup.get(key)
    if mapped:
        return mapped, "direct"

    stripped = TRAILING_STATE_PINCODE_RE.sub("", inst)
    if stripped != inst:
        mapped = lookup.get(normalize_institute(stripped))
        if mapped:
            return mapped, "suffix-stripped"

    return None, None


def clean_cell(x):
    """Collapse a table cell's internal line-wraps/whitespace into single spaces."""
    if x is None:
        return ""
    return " ".join(str(x).replace("\n", " ").split())


def is_row_start(cell):
    """A row's own SNo/Rank cell is a plain integer (commas allowed)."""
    if cell is None:
        return False
    return str(cell).strip().replace(",", "").isdigit()


# --------------------------------------------------------------------------
# Round-label detection (for nicer output headers only -- not needed for
# the matching logic itself)
# --------------------------------------------------------------------------
def detect_latest_round_label(pdf, probe_pages=8, top_cutoff=70):
    """
    Look at the first few pages for the header sub-line listing "Round 1
    Round 2 ... Round N" and return the LAST label found (e.g. "Round 3").
    Returns None if this PDF doesn't have that sub-header at all (i.e. the
    simple, Round-1-style layout).
    """
    total_pages = len(pdf.pages)
    for page_idx in range(min(probe_pages, total_pages)):
        page = pdf.pages[page_idx]
        words = page.extract_words(use_text_flow=False, keep_blank_chars=False)
        page.flush_cache()
        if not words:
            continue
        header_words = [w for w in words if w["top"] < top_cutoff]
        # Group into lines by 'top' to isolate the specific sub-header
        # line, not the page title (which also mentions one "Round N").
        by_top = {}
        for w in header_words:
            by_top.setdefault(round(w["top"], 0), []).append(w)
        best_labels = []
        for top, ws in by_top.items():
            ws_sorted = sorted(ws, key=lambda w: w["x0"])
            line_text = " ".join(w["text"] for w in ws_sorted)
            labels = ROUND_LABEL_RE.findall(line_text)
            if len(labels) > len(best_labels):
                best_labels = labels
        if len(best_labels) >= 2:
            return best_labels[-1]
    return None


# --------------------------------------------------------------------------
# Parsing -- uses pdfplumber's own table detection (handles any number of
# round-blocks and any page width automatically; no hardcoded pixel
# boundaries to maintain).
# --------------------------------------------------------------------------
def parse_records(pdf, skip_pages=2, progress_every=200):
    """
    Returns (records, layout) where layout is "simple" (8-column, one
    Allotted Institute column) or "dual" (12+ column, latest round's
    Allotted Institute is the 6th-from-last column). Each record is
    {"id": ..., "institute": ..., "other": ...} where "other" is the
    remaining columns for that same (latest) round, joined for reference.
    """
    records = []
    layout = None
    total_pages = len(pdf.pages)

    for page_idx in range(skip_pages, total_pages):
        page = pdf.pages[page_idx]
        try:
            tables = page.extract_tables()
        except Exception:
            tables = []
        for t in tables:
            for row in t:
                if row is None or len(row) < 8:
                    continue
                if not is_row_start(row[0]):
                    continue
                ncols = len(row)
                if layout is None:
                    layout = "simple" if ncols == 8 else "dual"
                if ncols == 8:
                    inst = row[3]
                    other = " | ".join(clean_cell(x) for x in row[4:] if x not in (None, ""))
                else:
                    inst = row[-6]
                    other = " | ".join(clean_cell(x) for x in row[-5:] if x not in (None, ""))
                records.append({
                    "id": clean_cell(row[0]),
                    "institute": clean_cell(inst),
                    "other": other,
                })
        page.flush_cache()
        if progress_every and (page_idx + 1) % progress_every == 0:
            print(f"  ...parsed page {page_idx + 1}/{total_pages}", file=sys.stderr, flush=True)

    return records, (layout or "simple")


# --------------------------------------------------------------------------
# FORMAT.xlsx lookup
# --------------------------------------------------------------------------
def load_format_lookup(format_xlsx_path, sheet_name=None):
    """
    Expected columns (header row 1): Allotted Institute, Code, Course,
    State, Institute Type, Institute Name
    """
    wb = openpyxl.load_workbook(format_xlsx_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]

    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_idx = {name: i for i, name in enumerate(header) if name}

    required = ["Allotted Institute", "Code", "State", "Institute Type", "Institute Name"]
    for r in required:
        if r not in col_idx:
            raise ValueError(f"FORMAT workbook is missing expected column '{r}'. Found columns: {header}")

    lookup = OrderedDict()
    for row in ws.iter_rows(min_row=2, values_only=True):
        inst_raw = row[col_idx["Allotted Institute"]]
        if not inst_raw:
            continue
        key = normalize_institute(inst_raw)
        lookup[key] = {
            "Code": row[col_idx["Code"]],
            "State": row[col_idx["State"]],
            "Institute Type": row[col_idx["Institute Type"]],
            "Institute Name": row[col_idx["Institute Name"]],
        }
    return lookup


# --------------------------------------------------------------------------
# Building output rows
# --------------------------------------------------------------------------
def build_output_rows(records, lookup, layout):
    """
    Simple layout: the one institute column is always used.
    Dual layout: only the latest round's own institute column is used --
    if it's blank/"-", the row is marked "N/A" and left unmapped, exactly
    as instructed (nothing pulled from an earlier round).

    Matching tries an exact normalized match first, then falls back to
    stripping a trailing ", State, Pincode" suffix (see lookup_institute)
    for PDFs whose address text includes that extra suffix. Rows matched
    only via that fallback are flagged in Match Status so it's always
    visible which rows needed it.
    """
    out_rows = []
    matched = 0
    matched_direct = 0
    matched_stripped = 0
    applicable = 0
    for rec in records:
        inst = rec["institute"].strip()
        has_institute = bool(inst) and inst != "-"

        if layout == "simple" or has_institute:
            if inst:
                applicable += 1
                mapped, via = lookup_institute(inst, lookup)
                if mapped:
                    matched += 1
                    if via == "direct":
                        matched_direct += 1
                        status = "Matched"
                    else:
                        matched_stripped += 1
                        status = "Matched (suffix-stripped)"
                else:
                    status = "No Match"
            else:
                mapped, status = None, "N/A"
        else:
            mapped, status = None, "N/A"

        mapped = mapped or {"Code": "", "State": "", "Institute Type": "", "Institute Name": ""}

        out_rows.append([
            rec["id"], inst, rec["other"],
            mapped["Code"], mapped["State"], mapped["Institute Type"], mapped["Institute Name"], status,
        ])
    stats = {
        "total": len(records), "matched": matched, "matched_direct": matched_direct,
        "matched_stripped": matched_stripped, "applicable": applicable,
    }
    return out_rows, stats


# --------------------------------------------------------------------------
# Writing the workbook
# --------------------------------------------------------------------------
def write_output_workbook(headers, out_rows, stats, out_path, sheet_title, unmatched_institutes=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    body_font = Font(name="Arial", size=10)
    no_match_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    status_col = headers.index("Match Status")
    ncols = len(headers)

    for row in out_rows:
        ws.append(row)

    for i, row in enumerate(out_rows, start=2):
        if row[status_col] == "No Match":
            for c in range(1, ncols + 1):
                cell = ws.cell(row=i, column=c)
                cell.font = body_font
                cell.fill = no_match_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"

    for i, name in enumerate(headers, start=1):
        if "Institute" in name and "reference" not in name.lower():
            ws.column_dimensions[get_column_letter(i)].width = 40
        elif "reference" in name.lower():
            ws.column_dimensions[get_column_letter(i)].width = 45
        elif name == "Institute Name":
            ws.column_dimensions[get_column_letter(i)].width = 32
        else:
            ws.column_dimensions[get_column_letter(i)].width = 14

    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2["A1"].font, ws2["B1"].font = header_font, header_font
    ws2["A1"].fill, ws2["B1"].fill = header_fill, header_fill
    rate = f"{stats['matched']/stats['applicable']*100:.1f}%" if stats["applicable"] else "N/A"
    for row in [
        ["Total rows extracted from PDF", stats["total"]],
        ["Rows with an institute to map (applicable)", stats["applicable"]],
        ["Rows marked N/A (blank/'-' in this round)", stats["total"] - stats["applicable"]],
        ["Matched to FORMAT lookup (total)", stats["matched"]],
        ["  - matched directly", stats["matched_direct"]],
        ["  - matched after stripping trailing State/Pincode", stats["matched_stripped"]],
        ["Match rate (of applicable rows)", rate],
    ]:
        ws2.append(row)
    for r in ws2.iter_rows(min_row=1, max_row=ws2.max_row):
        for cell in r:
            cell.font = body_font
    ws2.column_dimensions["A"].width = 48
    ws2.column_dimensions["B"].width = 20

    if unmatched_institutes:
        ws3 = wb.create_sheet("Unmatched Institutes (unique)")
        ws3.append(["Institute (raw text)", "Occurrences"])
        for cell in ws3[1]:
            cell.font, cell.fill = header_font, header_fill
        for inst, count in unmatched_institutes.items():
            ws3.append([inst, count])
        ws3.column_dimensions["A"].width = 90
        ws3.column_dimensions["B"].width = 14
        for r in ws3.iter_rows(min_row=2, max_row=ws3.max_row):
            for cell in r:
                cell.font = body_font

    wb.save(out_path)


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input_pdf", help="Path to ANY round's MCC allotment PDF (Round 1, 2, 3, 4, 5, ...)")
    parser.add_argument("output_xlsx", nargs="?", default=None,
                         help="Output Excel path. If omitted, derived from the PDF filename.")
    parser.add_argument("--format-xlsx", default=FORMAT_XLSX_PATH,
                         help=f"Path to the fixed FORMAT.xlsx lookup workbook (default: {FORMAT_XLSX_PATH})")
    parser.add_argument("--skip-pages", type=int, default=2,
                         help="Number of leading legend/cover pages to skip (default: 2)")
    args = parser.parse_args()

    output_xlsx = args.output_xlsx or str(Path(args.input_pdf).with_name(Path(args.input_pdf).stem + "_mapped.xlsx"))

    print(f"Loading FORMAT lookup from {args.format_xlsx} ...", file=sys.stderr)
    lookup = load_format_lookup(args.format_xlsx)
    print(f"  {len(lookup)} reference institutes loaded.", file=sys.stderr)

    print(f"Opening PDF {args.input_pdf} ...", file=sys.stderr)
    with pdfplumber.open(args.input_pdf) as pdf:
        latest_label = detect_latest_round_label(pdf)
        print("Parsing PDF (uses the PDF's own table structure -- works for any round) ...", file=sys.stderr)
        records, layout = parse_records(pdf, skip_pages=args.skip_pages)
        print(f"  Detected layout: {layout}" + (f" (latest = {latest_label})" if layout == "dual" else ""),
              file=sys.stderr)
        print(f"  {len(records)} rows extracted.", file=sys.stderr)

    out_rows, stats = build_output_rows(records, lookup, layout)

    if layout == "simple":
        inst_label = "Allotted Institute"
        other_label = "Course / Category / Remarks (reference only)"
        sheet_title = "Allotment Data"
    else:
        latest_tag = latest_label or "Latest Round"
        inst_label = f"[{latest_tag}] Allotted Institute"
        other_label = f"[{latest_tag}] Course / Category / Remarks (reference only)"
        sheet_title = f"{latest_tag} Allotment Data"

    headers = ["Rank", inst_label, other_label, "Code", "State", "Institute Type", "Institute Name", "Match Status"]

    unmatched_institutes = OrderedDict()
    status_idx = headers.index("Match Status")
    inst_idx = headers.index(inst_label)
    for row in out_rows:
        if row[status_idx] == "No Match":
            unmatched_institutes[row[inst_idx]] = unmatched_institutes.get(row[inst_idx], 0) + 1

    print(f"Writing output workbook to {output_xlsx} ...", file=sys.stderr)
    write_output_workbook(headers, out_rows, stats, output_xlsx, sheet_title, unmatched_institutes)
    print("Done.", file=sys.stderr)
    rate = f"{stats['matched']/stats['applicable']*100:.1f}%" if stats["applicable"] else "N/A"
    print(f"Matched {stats['matched']}/{stats['applicable']} applicable rows ({rate}). "
          f"{stats['total'] - stats['applicable']} rows were N/A.", file=sys.stderr)


if __name__ == "__main__":
    main()
